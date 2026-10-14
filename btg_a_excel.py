"""Convierte extractos PDF de BTG Pactual (iBanca) a Excel: Fecha | Concepto | Débito | Crédito.

Uso CLI:   python btg_a_excel.py extracto.pdf [otro.pdf ...]   -> genera <nombre>_extraido.xlsx
Uso API:   uvicorn btg_a_excel:app --reload   -> POST /convertir (archivo PDF) devuelve el .xlsx
"""
import io, re, sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

FECHA = re.compile(r"^\s*(\d{2}/\d{2}/\d{4})\s+(.*)$")
MONTO = re.compile(r"-?\d{1,3}(?:\.\d{3})*,\d{2}")


def a_decimal(txt):
    return Decimal(txt.replace(".", "").replace(",", "."))


def parsear(pdf_bytes):
    """Devuelve lista de dicts {fecha, concepto, debito, credito, saldo}."""
    filas = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        texto = "\n".join(p.extract_text(layout=True) or "" for p in pdf.pages)
    for linea in texto.splitlines():
        m = FECHA.match(linea)
        if not m:
            continue
        fecha, resto = m.groups()
        montos = list(MONTO.finditer(resto))
        if len(montos) < 2:
            continue
        importe, saldo = montos[-2].group(), montos[-1].group()
        descripcion = resto[: montos[-2].start()].strip()
        # Concepto y Referencia/Cheque vienen separados por 2+ espacios -> se unen con 1
        concepto = " ".join(p.strip() for p in re.split(r"\s{2,}", descripcion) if p.strip())
        valor = a_decimal(importe)
        filas.append({
            "fecha": datetime.strptime(fecha, "%d/%m/%Y"),
            "concepto": concepto,
            "debito": -valor if valor < 0 else Decimal("0"),
            "credito": valor if valor > 0 else Decimal("0"),
            "saldo": a_decimal(saldo),
        })
    return filas


def validar(filas):
    """Chequea que saldo anterior - débito + crédito = saldo de cada línea."""
    errores = []
    for i in range(1, len(filas)):
        esperado = filas[i - 1]["saldo"] - filas[i]["debito"] + filas[i]["credito"]
        if esperado != filas[i]["saldo"]:
            errores.append(f"Fila {i + 1}: saldo esperado {esperado}, PDF dice {filas[i]['saldo']}")
    return errores


def a_excel(filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Fecha", "Concepto", "Débito", "Crédito"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for f in filas:
        ws.append([f["fecha"], f["concepto"], float(f["debito"]), float(f["credito"])])
    for fila in ws.iter_rows(min_row=2):
        fila[0].number_format = "dd/mm/yy"
        fila[2].number_format = fila[3].number_format = "0.00"
    for col, ancho in {"A": 21, "B": 37, "C": 12, "D": 12}.items():
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def convertir(pdf_bytes):
    filas = parsear(pdf_bytes)
    return a_excel(filas), filas, validar(filas)


# ---------- API (opcional, requiere fastapi + uvicorn + python-multipart) ----------
try:
    from fastapi import FastAPI, File, HTTPException, UploadFile
    from fastapi.responses import Response

    app = FastAPI(title="BTG PDF a Excel")

    @app.post("/convertir")
    async def convertir_endpoint(archivo: UploadFile = File(...)):
        xlsx, filas, errores = convertir(await archivo.read())
        if not filas:
            raise HTTPException(422, "No se encontraron movimientos en el PDF")
        nombre = Path(archivo.filename or "extracto").stem + "_extraido.xlsx"
        return Response(
            xlsx,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{nombre}"',
                     "X-Movimientos": str(len(filas)),
                     "X-Errores-Saldo": str(len(errores))},
        )
except ImportError:
    app = None


if __name__ == "__main__":
    for ruta in sys.argv[1:]:
        p = Path(ruta)
        xlsx, filas, errores = convertir(p.read_bytes())
        salida = p.with_name(p.stem + "_extraido.xlsx")
        salida.write_bytes(xlsx)
        print(f"{p.name}: {len(filas)} movimientos -> {salida.name}")
        for e in errores:
            print("  ⚠", e)
